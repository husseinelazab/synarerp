"""استيراد البيانات وتصديرها لكل أنواع المستندات (Excel / CSV).

- الأعمدة تُشتق من حقول نموذج الإدخال نفسه بعناوين عربية، فالقالب يطابق الشاشة دائماً.
- الروابط (العميل، الصنف، الحساب...) تُطابق بالكود أو الرقم أو الاسم داخل الشركة الحالية فقط.
- المستندات ذات البنود: الأسطر التي تحمل نفس «المرجع» تُكوّن مستنداً واحداً.
- التحقق يمر بقواعد نموذج الإدخال نفسه، والمعاينة تنفذ الاستيراد كاملاً داخل معاملة ثم تتراجع عنه.
"""
import csv
import datetime
import io
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.db import IntegrityError, models, transaction
from django.forms import modelform_factory
from django.forms.models import model_to_dict

from core.forms import build_form_class, formfield_callback

REF_LABEL = "المرجع"
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩٫٬", "0123456789.,")
TRUE_WORDS = {"1", "true", "yes", "y", "نعم", "صح", "✓", "x"}
FALSE_WORDS = {"0", "false", "no", "n", "لا", "خطأ", ""}
# حقول تُستخدم لمطابقة السجلات المرتبطة، بالترتيب
LOOKUP_FIELDS = ("code", "number", "employee_number", "name", "full_name", "prefix")


@dataclass
class Column:
    key: str            # اسم الحقل، أو "__ref__"
    label: str
    field: object = None  # حقل النموذج (Model field)
    child: bool = False
    required: bool = False


class Rollback(Exception):
    pass


# ---------------- الأعمدة ----------------

def import_key(model):
    """الحقل الفريد داخل الشركة الذي يُستخدم لتحديث السجلات الموجودة."""
    for constraint in model._meta.constraints:
        fields = list(getattr(constraint, "fields", []))
        if len(fields) == 2 and fields[0] == "company":
            return fields[1]
    return None


def columns_for(dt):
    columns = []
    if dt.child:
        columns.append(Column("__ref__", REF_LABEL, required=True))
    form_class = build_form_class(dt)
    for name, form_field in form_class.base_fields.items():
        field = dt.model._meta.get_field(name)
        # الحقل الذي له قيمة افتراضية اختياري في الملف وتُستخدم قيمته الافتراضية
        required = form_field.required and not field.has_default() and not isinstance(field, models.BooleanField)
        columns.append(Column(name, str(field.verbose_name), field, required=required))
    if dt.child:
        for name in dt.child.fields:
            field = dt.child.model._meta.get_field(name)
            columns.append(Column(name, f"{dt.child.label}: {field.verbose_name}", field, child=True,
                                  required=not field.blank and not field.has_default()))
    return columns


# ---------------- تحويل القيم ----------------

def _clean_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def resolve_fk(field, raw):
    model = field.related_model
    queryset = model._default_manager.complex_filter(field.get_limit_choices_to() or {})
    text = _clean_text(raw).translate(DIGITS)
    candidates = [text]
    if " - " in text:  # صيغة «الكود - الاسم» كما تظهر في القوائم والتصدير
        candidates += [part.strip() for part in text.split(" - ", 1)]
    lookups = [f for f in LOOKUP_FIELDS if any(x.name == f for x in model._meta.concrete_fields)]
    for candidate in candidates:
        for lookup in lookups:
            matches = list(queryset.filter(**{f"{lookup}__iexact": candidate})[:2])
            if len(matches) == 1:
                return matches[0]
            if len(matches) > 1:
                raise ValidationError(f"القيمة «{text}» تطابق أكثر من سجل في «{model._meta.verbose_name_plural}».")
    raise ValidationError(f"لم يتم العثور على «{text}» في «{model._meta.verbose_name_plural}».")


def parse_value(field, raw):
    """تحويل قيمة الخلية إلى قيمة يقبلها نموذج الإدخال."""
    if isinstance(field, models.BooleanField):
        text = _clean_text(raw).lower()
        if isinstance(raw, bool):
            return raw
        if text in TRUE_WORDS:
            return True
        if text in FALSE_WORDS:
            return False
        raise ValidationError(f"القيمة «{raw}» غير مفهومة؛ استخدم نعم أو لا.")
    if raw is None or _clean_text(raw) == "":
        return None
    if isinstance(field, models.ForeignKey):
        return resolve_fk(field, raw).pk
    if field.choices:
        text = _clean_text(raw)
        for value, label in field.flatchoices:
            if text in (str(value), str(label)):
                return value
        allowed = "، ".join(str(label) for _, label in field.flatchoices)
        raise ValidationError(f"القيمة «{text}» غير مسموحة. القيم المسموحة: {allowed}.")
    if isinstance(field, models.DateField):
        if isinstance(raw, datetime.datetime):
            return raw.date()
        if isinstance(raw, datetime.date):
            return raw
        text = _clean_text(raw).translate(DIGITS)
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        raise ValidationError(f"التاريخ «{text}» غير صحيح؛ استخدم الصيغة 2026-01-31.")
    if isinstance(field, (models.DecimalField, models.IntegerField, models.FloatField)):
        text = _clean_text(raw).translate(DIGITS).replace(",", "")
        try:
            number = Decimal(text)
        except InvalidOperation:
            raise ValidationError(f"«{raw}» ليس رقماً.")
        return int(number) if isinstance(field, models.IntegerField) else number
    return _clean_text(raw)


# ---------------- قراءة الملفات ----------------

def read_table(filename, content):
    """تعيد (العناوين، الصفوف) من ملف xlsx أو csv."""
    if filename.lower().endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1256"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValidationError("تعذر قراءة ملف CSV؛ احفظه بترميز UTF-8.")
        rows = list(csv.reader(io.StringIO(text)))
    else:
        import openpyxl

        try:
            workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception:
            raise ValidationError("الملف ليس ملف Excel صالحاً (xlsx) أو CSV.")
        sheet = workbook["البيانات"] if "البيانات" in workbook.sheetnames else workbook.worksheets[0]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    rows = [row for row in rows if any(_clean_text(cell) for cell in row)]
    if not rows:
        raise ValidationError("الملف فارغ.")
    headers = [_clean_text(h).rstrip("*").strip() for h in rows[0]]
    return headers, rows[1:]


def map_headers(dt, headers, mode="insert"):
    columns = columns_for(dt)
    by_label = {}
    for column in columns:
        by_label[column.label] = column
        by_label.setdefault(column.key, column)
    mapping, unknown = {}, []
    for index, header in enumerate(headers):
        if not header:
            continue
        column = by_label.get(header)
        if column is None:
            unknown.append(header)
        else:
            mapping[index] = column
    key = import_key(dt.model)
    if mode == "update" and not dt.is_submittable and key:
        # التحديث يكفيه الحقل الفريد (مثل كود الصنف) مع الأعمدة المراد تغييرها فقط؛ الباقي يبقى كما هو
        missing = [c.label for c in columns if c.key == key and c not in mapping.values()]
    else:
        missing = [c.label for c in columns if c.required and c not in mapping.values()]
    return mapping, unknown, missing


# ---------------- التنفيذ ----------------

def _group_rows(dt, mapping, rows):
    """تجميع أسطر المستند الواحد بالمرجع؛ السطر بلا مرجع يتبع ما قبله."""
    if not dt.child:
        return [[(i + 2, row)] for i, row in enumerate(rows)]
    ref_index = next((i for i, c in mapping.items() if c.key == "__ref__"), None)
    groups, by_ref, current = [], {}, None
    for i, row in enumerate(rows):
        ref = _clean_text(row[ref_index]) if ref_index is not None and ref_index < len(row) else ""
        if ref:
            if ref not in by_ref:
                by_ref[ref] = []
                groups.append(by_ref[ref])
            current = by_ref[ref]
        elif current is None:
            current = []
            groups.append(current)
        current.append((i + 2, row))
    return groups


def _form_errors(form, columns_by_key, line, skip=()):
    messages = []
    for name, errors in form.errors.items():
        if name in skip:  # الحقل له خطأ قراءة سابق؛ لا نكرر الرسالة
            continue
        label = columns_by_key[name].label if name in columns_by_key else ""
        for error in errors:
            messages.append({"row": line, "column": label, "message": error})
    return messages


def _cells(mapping, row, child):
    values = {}
    for index, column in mapping.items():
        if column.key == "__ref__" or column.child != child:
            continue
        values[column.key] = (column, row[index] if index < len(row) else None)
    return values


def _build_data(base, cells, line):
    data, errors = dict(base), []
    for key, (column, raw) in cells.items():
        try:
            value = parse_value(column.field, raw)
        except ValidationError as exc:
            errors.append({"row": line, "column": column.label, "message": " ".join(exc.messages)})
            continue
        data[key] = "" if value is None else value
    return data, errors


def _failed_fields(errors, cells):
    labels = {e["column"] for e in errors}
    return {key for key, (column, _) in cells.items() if column.label in labels}


def _import_group(dt, mapping, group, mode, submit, user):
    form_class = build_form_class(dt)
    columns_by_key = {c.key: c for c in mapping.values()}
    first_line, first_row = group[0]
    parent_cells = _cells(mapping, first_row, child=False)

    instance, created = None, True
    key = import_key(dt.model)
    if mode == "update" and not dt.is_submittable and key and key in parent_cells:
        value = _clean_text(parent_cells[key][1])
        instance = dt.model._default_manager.filter(**{key: value}).first()
        created = instance is None
    base_obj = instance or dt.model()
    base = {k: v for k, v in model_to_dict(base_obj, fields=list(form_class.base_fields)).items()}
    data, errors = _build_data(base, parent_cells, first_line)
    form = form_class(data=data, instance=instance)
    if errors or not form.is_valid():
        # نجمع أخطاء القراءة وأخطاء التحقق معاً ليصحح المستخدم الصف دفعة واحدة
        raise GroupError(errors + _form_errors(form, columns_by_key, first_line,
                                               skip=_failed_fields(errors, parent_cells)))
    obj = form.save(commit=False)
    if hasattr(obj, "created_by_id") and created:
        obj.created_by = user
    obj.save()

    if dt.child:
        child_form_class = modelform_factory(dt.child.model, fields=list(dt.child.fields),
                                             formfield_callback=formfield_callback)
        child_errors = []
        for line, row in group:
            cells = _cells(mapping, row, child=True)
            if not any(_clean_text(raw) for _, raw in cells.values()):
                continue
            row_data, errors = _build_data(model_to_dict(dt.child.model(), fields=list(dt.child.fields)), cells, line)
            child_form = child_form_class(data=row_data)
            if errors or not child_form.is_valid():
                child_errors += errors + _form_errors(child_form, columns_by_key, line,
                                                      skip=_failed_fields(errors, cells))
                continue
            child = child_form.save(commit=False)
            setattr(child, dt.child.fk_name, obj)
            child.save()
        if child_errors:
            raise GroupError(child_errors)

    if dt.is_submittable:
        obj.calculate_totals()
        obj.save()
        if submit:
            obj.submit()
    return obj, created


class GroupError(Exception):
    def __init__(self, errors):
        super().__init__("group failed")
        self.errors = errors


def run_import(dt, filename, content, mode="insert", submit=False, valid_only=False, dry_run=True, user=None):
    result = {"created": 0, "updated": 0, "failed": 0, "total": 0, "errors": [], "unknown_columns": [],
              "missing_columns": [], "documents": []}
    try:
        headers, rows = read_table(filename, content)
    except ValidationError as exc:
        result["errors"].append({"row": "", "column": "", "message": " ".join(exc.messages)})
        return result
    mapping, unknown, missing = map_headers(dt, headers, mode)
    result["unknown_columns"], result["missing_columns"] = unknown, missing
    if missing:
        result["errors"].append({"row": 1, "column": "", "message": "أعمدة مطلوبة غير موجودة: " + "، ".join(missing)})
        return result

    groups = _group_rows(dt, mapping, rows)
    result["total"] = len(groups)
    try:
        with transaction.atomic():
            for group in groups:
                savepoint = transaction.savepoint()
                try:
                    from stock.pricing import price_source

                    with price_source("import", filename):
                        obj, created = _import_group(dt, mapping, group, mode, submit, user)
                except GroupError as exc:
                    transaction.savepoint_rollback(savepoint)
                    result["failed"] += 1
                    result["errors"] += exc.errors
                except (ValidationError, IntegrityError) as exc:
                    transaction.savepoint_rollback(savepoint)
                    result["failed"] += 1
                    message = " ".join(exc.messages) if hasattr(exc, "messages") else "يوجد سجل آخر بنفس القيمة في هذه الشركة."
                    result["errors"].append({"row": group[0][0], "column": "", "message": message})
                else:
                    transaction.savepoint_commit(savepoint)
                    result["created" if created else "updated"] += 1
                    if len(result["documents"]) < 200:
                        result["documents"].append({"pk": obj.pk, "label": str(obj)})
            if dry_run or (result["failed"] and not valid_only):
                raise Rollback
    except Rollback:
        if not dry_run:
            result["created"] = result["updated"] = 0
            result["documents"] = []
    return result


# ---------------- القوالب والتصدير ----------------

def _cell_value(obj, column):
    value = getattr(obj, column.key)
    if value is None:
        return ""
    if column.field.choices:
        return getattr(obj, f"get_{column.key}_display")()
    if isinstance(column.field, models.BooleanField):
        return "نعم" if value else "لا"
    if isinstance(column.field, models.ForeignKey):
        return str(value)
    if isinstance(value, Decimal):
        return float(value)
    return value


def export_rows(dt, queryset):
    columns = columns_for(dt)
    rows = []
    for obj in queryset:
        parent = [obj.name if c.key == "__ref__" else _cell_value(obj, c) for c in columns if not c.child]
        if dt.child:
            related = getattr(obj, dt.child.model._meta.get_field(dt.child.fk_name).remote_field.related_name).all()
            for i, line in enumerate(related):
                child = [_cell_value(line, c) for c in columns if c.child]
                rows.append((parent if i == 0 else [obj.name] + [""] * (len(parent) - 1)) + child)
            if not related:
                rows.append(parent + [""] * len([c for c in columns if c.child]))
        else:
            rows.append(parent)
    return columns, rows


def build_workbook(dt, columns, rows, with_help=True):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "البيانات"
    sheet.sheet_view.rightToLeft = True
    sheet.append([c.label + (" *" if c.required else "") for c in columns])
    header_fill, required_fill = PatternFill("solid", fgColor="1F4E79"), PatternFill("solid", fgColor="8B1F2A")
    for index, column in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if column.required else header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[cell.column_letter].width = max(14, min(40, len(cell.value) + 6))
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(row)

    if with_help:
        help_sheet = workbook.create_sheet("تعليمات")
        help_sheet.sheet_view.rightToLeft = True
        help_sheet.append(["العمود", "مطلوب", "النوع / القيم المسموحة"])
        for cell in help_sheet[1]:
            cell.font = Font(bold=True)
        for column in columns:
            help_sheet.append([column.label, "نعم" if column.required else "", describe(column)])
        help_sheet.append([])
        notes = [
            "الأعمدة الحمراء مطلوبة. يمكن حذف الأعمدة غير المطلوبة التي لا تحتاجها.",
            "الروابط (عميل، صنف، حساب...) تُكتب بالكود أو الرقم أو الاسم كما يظهر في النظام.",
            "التواريخ بالصيغة 2026-01-31، والأرقام بدون رمز العملة.",
        ]
        if dt.child:
            notes.insert(0, f"كل الأسطر التي تحمل نفس «{REF_LABEL}» تُنشئ مستنداً واحداً؛ "
                            "تكفي بيانات المستند في أول سطر، والأسطر التالية للبنود فقط.")
        for note in notes:
            help_sheet.append([note])
        help_sheet.column_dimensions["A"].width = 40
        help_sheet.column_dimensions["C"].width = 90
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def describe(column):
    field = column.field
    if column.key == "__ref__":
        return "رقم أو نص يجمع أسطر المستند الواحد (مثل 1، 2، 3)"
    if field.choices:
        return "إحدى القيم: " + "، ".join(str(label) for _, label in field.flatchoices)
    if isinstance(field, models.ForeignKey):
        examples = [str(o) for o in field.related_model._default_manager.complex_filter(
            field.get_limit_choices_to() or {})[:5]]
        return f"من «{field.related_model._meta.verbose_name_plural}» مثل: " + "، ".join(examples)
    if isinstance(field, models.BooleanField):
        return "نعم / لا"
    if isinstance(field, models.DateField):
        return "تاريخ: 2026-01-31"
    if isinstance(field, (models.DecimalField, models.IntegerField)):
        return "رقم"
    return "نص"


def build_csv(columns, rows):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([c.label + (" *" if c.required else "") for c in columns])
    writer.writerows(rows)
    return ("﻿" + output.getvalue()).encode("utf-8")
