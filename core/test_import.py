import io
from decimal import Decimal

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import Sum
from django.test import TestCase

from accounts.models import Account, GLEntry, JournalEntry
from core.importer import columns_for, run_import
from core.models import Company, DataImport, Membership, Role
from core.registry import REGISTRY, get_doctype
from saas.tests import make_company
from selling.models import Customer, SalesInvoice
from stock.models import Item


def xlsx(rows):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


class ImportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.company, cls.user = make_company("شركة الاستيراد", "importer")

    def setUp(self):
        self.client.force_login(self.user)

    def upload(self, slug, rows, filename="data.xlsx", **options):
        content = xlsx(rows) if filename.endswith(".xlsx") else rows
        data = {"file": SimpleUploadedFile(filename, content), "mode": options.get("mode", "insert")}
        if options.get("submit_after"):
            data["submit_after"] = "on"
        if options.get("valid_only"):
            data["valid_only"] = "on"
        response = self.client.post(f"/app/{slug}/import/", data)
        job = DataImport.all_objects.latest("id")
        self.assertRedirects(response, f"/app/{slug}/import/{job.pk}/")
        return job

    def execute(self, slug, job):
        self.client.post(f"/app/{slug}/import/{job.pk}/")
        job.refresh_from_db()
        return job

    def test_every_doctype_has_downloadable_template(self):
        from core.permissions import NONPROFIT_MODULES

        for slug, dt in REGISTRY.items():
            if dt.module in NONPROFIT_MODULES:  # وحدات الجمعيات مغلقة على المنشأة التجارية
                continue
            response = self.client.get(f"/app/{slug}/template.xlsx")
            self.assertEqual(response.status_code, 200, slug)
            book = openpyxl.load_workbook(io.BytesIO(response.content))
            self.assertEqual(book.sheetnames, ["البيانات", "تعليمات"])
            self.assertTrue(book["البيانات"].sheet_view.rightToLeft)
            self.assertEqual(self.client.get(f"/app/{slug}/template.csv").status_code, 200)
            self.assertEqual(self.client.get(f"/app/{slug}/import/").status_code, 200)

    def test_preview_then_execute_customers_and_update(self):
        rows = [["اسم العميل *", "نوع العميل", "الجوال", "المدينة"],
                ["مؤسسة الواحة", "شركة / منشأة", "0501111111", "الرياض"],
                ["سالم العتيبي", "فرد", "0502222222", "جدة"]]
        job = self.upload("customer", rows)
        self.assertEqual(job.result["created"], 2)
        self.assertFalse(Customer.objects.filter(name="مؤسسة الواحة").exists())  # المعاينة لا تحفظ شيئاً
        job = self.execute("customer", job)
        self.assertEqual(job.status, "done")
        self.assertEqual(Customer.objects.get(name="سالم العتيبي").customer_type, "individual")
        # التحديث بالاسم
        job = self.execute("customer", self.upload("customer", [["اسم العميل", "الجوال"], ["سالم العتيبي", "0509999999"]],
                                                   mode="update"))
        self.assertEqual(job.result["updated"], 1)
        customer = Customer.objects.get(name="سالم العتيبي")
        self.assertEqual(customer.phone, "0509999999")
        self.assertEqual(customer.city, "جدة")  # الأعمدة غير الموجودة في الملف لا تُمسح

    def test_items_resolve_links_by_name(self):
        rows = [["كود الصنف", "اسم الصنف", "مجموعة الأصناف", "وحدة القياس", "صنف مخزني", "سعر البيع"],
                ["A-1", "قهوة", "المنتجات", "كيلوجرام", "نعم", "45.5"],
                ["S-1", "توصيل", "الخدمات", "خدمة", "لا", "١٥"]]
        self.execute("item", self.upload("item", rows))
        self.assertEqual(Item.objects.get(code="A-1").item_group.name, "المنتجات")
        service = Item.objects.get(code="S-1")
        self.assertFalse(service.is_stock_item)
        self.assertEqual(service.selling_rate, Decimal("15"))

    def test_invoices_with_lines_grouped_by_reference_and_posted(self):
        Item.objects.create(code="SV", name="خدمة استشارية", item_group_id=Item._meta.get_field("item_group")
                            .related_model.objects.first().pk, uom_id=Item._meta.get_field("uom").related_model
                            .objects.first().pk, is_stock_item=False, selling_rate=100)
        Customer.objects.create(name="عميل الاستيراد")
        rows = [["المرجع", "العميل", "التاريخ", "تحديث المخزون", "البنود: الصنف", "البنود: الكمية", "البنود: السعر"],
                ["1", "عميل الاستيراد", "2026-03-01", "لا", "SV", "2", "100"],
                ["1", "", "", "", "خدمة استشارية", "1", "50"],
                ["2", "عميل الاستيراد", "2026-03-02", "لا", "SV", "1", "200"]]
        job = self.execute("sales-invoice", self.upload("sales-invoice", rows, submit_after=True))
        self.assertEqual(job.result["created"], 2, job.result["errors"])
        invoices = SalesInvoice.objects.order_by("posting_date")
        self.assertEqual([inv.items.count() for inv in invoices], [2, 1])
        self.assertEqual(invoices[0].grand_total, Decimal("287.50"))
        self.assertTrue(all(inv.is_submitted for inv in invoices))
        agg = GLEntry.objects.filter(is_cancelled=False).aggregate(d=Sum("debit"), c=Sum("credit"))
        self.assertEqual(agg["d"], agg["c"])

    def test_errors_reported_with_row_numbers_and_all_or_nothing(self):
        rows = [["اسم العميل", "نوع العميل"], ["صحيح", "فرد"], ["", "فرد"], ["خطأ النوع", "مجهول"]]
        job = self.upload("customer", rows)
        self.assertEqual(job.result["failed"], 2)
        rows_with_errors = {e["row"] for e in job.result["errors"]}
        self.assertEqual(rows_with_errors, {3, 4})
        self.assertIn("القيم المسموحة", " ".join(e["message"] for e in job.result["errors"]))
        # كل أخطاء الصف الواحد تظهر معاً
        both = self.upload("customer", [["اسم العميل", "حد الائتمان"], ["", "abc"]])
        self.assertEqual({e["column"] for e in both.result["errors"]}, {"اسم العميل", "حد الائتمان"})
        job = self.execute("customer", job)  # لا يُنفذ مع وجود أخطاء
        self.assertFalse(Customer.objects.filter(name="صحيح").exists())
        job = self.execute("customer", self.upload("customer", rows, valid_only=True))
        self.assertTrue(Customer.objects.filter(name="صحيح").exists())
        self.assertEqual(job.result["created"], 1)

    def test_duplicate_within_company_is_rejected(self):
        Customer.objects.create(name="مكرر")
        job = self.upload("customer", [["اسم العميل"], ["مكرر"]])
        self.assertEqual(job.result["failed"], 1)

    def test_csv_with_arabic_digits_and_journal_lines(self):
        csv_text = ("المرجع,التاريخ,نوع القيد,الحسابات: الحساب,الحسابات: مدين,الحسابات: دائن\n"
                    "ق1,2026-02-01,قيد افتتاحي,1101,١٠٠٠,0\n"
                    "ق1,,,3101 - رأس المال,0,1000\n").encode("utf-8-sig")
        job = self.execute("journal-entry", self.upload("journal-entry", csv_text, filename="je.csv", submit_after=True))
        self.assertEqual(job.result["created"], 1, job.result["errors"])
        entry = JournalEntry.objects.get()
        self.assertTrue(entry.is_submitted)
        self.assertEqual(entry.total_debit, Decimal("1000"))

    def test_accounts_import_with_parent(self):
        rows = [["رقم الحساب", "اسم الحساب", "الحساب الأب", "النوع الجذري", "نوع الحساب"],
                ["1199", "عهد الموظفين", "11", "أصول", ""]]
        job = self.execute("account", self.upload("account", rows))
        self.assertEqual(job.result["created"], 1, job.result["errors"])
        self.assertEqual(Account.objects.get(number="1199").parent.number, "11")

    def test_export_roundtrip_and_tenant_isolation(self):
        Customer.objects.create(name="عميل للتصدير", city="مكة")
        response = self.client.get("/app/customer/export/")
        self.assertEqual(response.status_code, 200)
        sheet = openpyxl.load_workbook(io.BytesIO(response.content))["البيانات"]
        names = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
        self.assertIn("عميل للتصدير", names)
        # شركة أخرى لا ترى بيانات الأولى ولا تستطيع الربط بها
        other, other_user = make_company("شركة أخرى", "other_importer")
        self.client.force_login(other_user)
        self.assertNotIn("عميل للتصدير", self.client.get("/app/customer/export/").content.decode("latin-1"))
        rows = [["المرجع", "العميل", "البنود: الصنف", "البنود: الكمية"], ["1", "عميل للتصدير", "X", "1"]]
        job = self.upload("sales-invoice", rows)
        self.assertEqual(job.result["failed"], 1)
        self.assertIn("لم يتم العثور", job.result["errors"][0]["message"])

    def test_columns_follow_form_fields(self):
        labels = [c.label for c in columns_for(get_doctype("sales-invoice"))]
        self.assertEqual(labels[0], "المرجع")
        self.assertIn("البنود: الكمية", labels)
