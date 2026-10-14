"""حزمة القوائم المالية: صفحة للطباعة أو الحفظ PDF، وملف Excel للمراجع."""
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import render

from core.models import FiscalYear


def _fiscal_year(request):
    years = list(FiscalYear.objects.order_by("-start_date"))
    selected = next((y for y in years if str(y.pk) == request.GET.get("year")), None)
    return years, selected or (years[0] if years else None)


@login_required
def statements_package(request):
    from core.models import PrintSettings
    from funds.statements import CLASSES, build_package

    years, year = _fiscal_year(request)
    if year is None:
        return render(request, "funds/package.html", {"years": years})
    package = build_package(request.company, year)
    return render(request, "funds/package.html", {
        **package, "years": years, "classes": CLASSES, "ps": PrintSettings.for_company(request.company)})


@login_required
def statements_export(request):
    import io

    import openpyxl
    from openpyxl.styles import Alignment, Font

    from funds.statements import CLASSES, build_package

    _, year = _fiscal_year(request)
    if year is None:
        return HttpResponse("لا توجد سنة مالية.", status=404)
    p = build_package(request.company, year)
    cur, pri = f"{p['end']:%Y-%m-%d}", f"{p['prior_end']:%Y-%m-%d}"
    book = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(title, header):
        ws = book.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append([request.company.name])
        ws.append([title])
        ws.append(header)
        for row in ws.iter_rows(min_row=1, max_row=3):
            for cell in row:
                cell.font = bold
        ws.column_dimensions["A"].width = 55
        for col in "BCDEFGHIJ":
            ws.column_dimensions[col].width = 18
        return ws

    def num(value):
        return float(value) if value is not None else None

    ws = sheet("قائمة المركز المالي", ["البند", "إيضاح", cur, pri])
    pos = p["position"]
    for group_title, groups in (("الموجودات", pos["assets"]), ("المطلوبات", pos["liabilities"])):
        ws.append([group_title])
        for section in groups:
            ws.append([section["title"]])
            for line in section["lines"]:
                ws.append([line["label"], line.get("note"), num(line["value"]), num(line["prior"])])
            ws.append([f"مجموع {section['title']}", None, num(section["value"]), num(section["prior"])])
    ws.append(["صافي الأصول"])
    for item in pos["net_assets"]:
        ws.append([item["label"], item.get("note"), num(item["value"]), num(item["prior"])])
    ws.append(["إجمالي صافي الأصول", None, num(pos["total_net"]), num(pos["total_net_prior"])])
    ws.append(["إجمالي المطلوبات وصافي الأصول", None, num(pos["total_liabilities"] + pos["total_net"]),
               num(pos["total_liabilities_prior"] + pos["total_net_prior"])])

    act = p["activities"]
    ws = sheet("قائمة الأنشطة", ["البيان", "إيضاح"] + [label for _, label in CLASSES] + [cur, pri])
    for section in act["sections"]:
        ws.append([section["title"]])
        for line in section["lines"]:
            ws.append([line["label"], line.get("note")] + [num(v) for v in line["values"].values()]
                      + [num(line["total"]), num(line["prior"])])
        ws.append([f"إجمالي {section['title']}", None] + [num(v) for v in section["totals"].values()]
                  + [num(section["total"]), num(section["prior"])])
    for t in act["transfers"]:
        ws.append([t["label"], None] + [num(v) for v in t["values"].values()] + [num(t["total"]), num(t["prior"])])
    for label, values, total, prior in (
            ("التغير في صافي الأصول من الأنشطة", act["change"], act["change_total"], act["change_prior"]),
            ("قيود مباشرة على صافي الأصول", act["direct"], act["direct_total"], act["direct_prior"]),
            ("صافي الأصول في بداية السنة", act["opening"], act["opening_total"], act["opening_prior"]),
            ("صافي الأصول في نهاية السنة", act["closing"], act["closing_total"], act["closing_prior"])):
        ws.append([label, None] + [num(v) for v in values.values()] + [num(total), num(prior)])

    ws = sheet("قائمة التدفقات النقدية", ["البيان", cur, pri])
    f, fp = p["flows"], p["flows_prior"]

    def prior_of(lines, label):
        return next((num(l["value"]) for l in lines if l["label"] == label), None)

    ws.append(["التدفقات النقدية من الأنشطة التشغيلية"])
    ws.append(["التغير في صافي الأصول", num(f["change_in_net_assets"]), num(fp["change_in_net_assets"])])
    for key in ("adjustments", "working"):
        for line in f[key]:
            ws.append([line["label"], num(line["value"]), prior_of(fp[key], line["label"])])
    ws.append(["صافي التدفقات النقدية من الأنشطة التشغيلية", num(f["operating"]), num(fp["operating"])])
    ws.append(["التدفقات النقدية من الأنشطة الاستثمارية"])
    for line in f["investing_lines"]:
        ws.append([line["label"], num(line["value"]), prior_of(fp["investing_lines"], line["label"])])
    ws.append(["صافي التدفقات النقدية من الأنشطة الاستثمارية", num(f["investing"]), num(fp["investing"])])
    ws.append(["التدفقات النقدية من الأنشطة التمويلية"])
    for line in f["financing_lines"]:
        ws.append([line["label"], num(line["value"]), prior_of(fp["financing_lines"], line["label"])])
    ws.append(["صافي التدفقات النقدية من الأنشطة التمويلية", num(f["financing"]), num(fp["financing"])])
    ws.append(["التغير في النقدية خلال السنة", num(f["net"]), num(fp["net"])])
    ws.append(["رصيد النقدية وما في حكمها بداية السنة", num(f["opening"]), num(fp["opening"])])
    ws.append(["رصيد النقدية وما في حكمها نهاية السنة", num(f["closing"]), num(fp["closing"])])

    ws = sheet("الإيضاحات", ["البند", cur, pri])
    for note in p["notes"]:
        ws.append([f"({note['number']}) {note['title']}"])
        ws.cell(ws.max_row, 1).font = bold
        if note["kind"] == "text":
            for paragraph in note["text"].splitlines():
                ws.append([paragraph])
                ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True)
        if note["kind"] == "fixed_assets":
            ws.append(["الفئة", "التكلفة أول السنة", "الإضافات", "الاستبعادات", "التكلفة آخر السنة", "مجمع الإهلاك أول السنة",
                       "إهلاك السنة", "إهلاك المستبعد", "مجمع الإهلاك آخر السنة", "الصافي"])
            for row in note["schedule"]["rows"]:
                ws.append([row["label"]] + [num(row[k]) for k in ("cost_open", "additions", "disposals", "cost_close",
                                                                 "dep_open", "dep_charge", "dep_disposals",
                                                                 "dep_close", "nbv")])
        if note["kind"] == "provision":
            m = note["movement"]
            for label, key in (("الرصيد أول العام", "opening"), ("المكون خلال العام", "additions"),
                               ("المستخدم خلال العام", "used"), ("الرصيد آخر العام", "closing")):
                ws.append([label, num(m[key])])
        for row in note.get("rows") or []:
            ws.append([row["label"], num(row["value"]), num(row["prior"])])
        if note.get("rows"):
            ws.append(["الإجمالي", num(note["total"]), num(note["total_prior"])])
        ws.append([])
    book.remove(book["Sheet"])
    for ws in book.worksheets:
        for row in ws.iter_rows(min_row=4):
            for cell in row[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = '#,##0.00;(#,##0.00);"-"'
    buffer = io.BytesIO()
    book.save(buffer)
    response = HttpResponse(buffer.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="financial-statements-{year.name}.xlsx"'
    return response
